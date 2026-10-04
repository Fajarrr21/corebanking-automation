"""build-bugs.py — generate bugs.xlsx dari bugs.csv (sumber kebenaran).

  python scripts/build-bugs.py

bugs.xlsx hanya hasil generate — JANGAN diedit manual. Tambah/ubah bug di bugs.csv
(ID berurutan BUG-###).
"""
import csv
import pathlib
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
COLS = ["Bug ID", "Module", "Title", "Steps to Reproduce", "Expected Result",
        "Actual Result", "Severity", "Priority", "Reporter", "Status", "Date Found"]
WIDTHS = [10, 18, 40, 50, 40, 40, 10, 10, 14, 12, 12]

rows = list(csv.DictReader((ROOT / "bugs.csv").open(encoding="utf-8")))
ids = [r["Bug ID"] for r in rows]
for i, bid in enumerate(ids, 1):
    if bid != f"BUG-{i:03d}":
        raise SystemExit(f"ID tidak berurutan: baris {i} = {bid}, seharusnya BUG-{i:03d}")

wb = Workbook()
ws = wb.active
ws.title = "Bugs"
for c, (name, w) in enumerate(zip(COLS, WIDTHS), 1):
    cell = ws.cell(row=1, column=c, value=name)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F3864")
    ws.column_dimensions[cell.column_letter].width = w
for r, row in enumerate(rows, 2):
    for c, name in enumerate(COLS, 1):
        ws.cell(row=r, column=c, value=row.get(name, "")).alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"
wb.save(ROOT / "bugs.xlsx")
print(f"OK bugs.xlsx ({len(rows)} bug)")
