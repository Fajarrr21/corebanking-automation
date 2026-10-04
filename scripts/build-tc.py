"""build-tc.py — generate TC sheet xlsx dari sumber JSON.

Sumber kebenaran TC = docs/test-cases/src/TC_<Modul>_<Aksi>.json
Hasil            = docs/test-cases/TC_<Modul>_<Aksi>.xlsx  (JANGAN edit manual)

Pemakaian:
  python scripts/build-tc.py            # build semua
  python scripts/build-tc.py TC_CIF_List  # build satu (nama file tanpa .json)

Format JSON:
{
  "modul": "Nasabah - CIF", "kode": "CIF", "aksi": "LIST",
  "halaman": "/cif",
  "notion": [{"judul": "Data CIF Nasabah", "url": "https://..."}],
  "catatan": "opsional",
  "cases": [{
    "id": "TC-CIF-LIST-001", "kategori": "Happy|Positif|Negatif|Edge",
    "skenario": "...", "precondition": "...",
    "langkah": ["...", "..."] | "...", "test_data": "...", "expected": "...",
    "sumber": "PRD|PRD-ambigu|Asumsi", "status": "Belum dijalankan", "catatan": "..."
  }]
}
"""
import json
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "test-cases" / "src"
OUT = ROOT / "docs" / "test-cases"

COLUMNS = [
    ("ID", 20), ("Kategori", 11), ("Skenario", 40), ("Pre-condition", 32),
    ("Langkah", 50), ("Test Data", 28), ("Expected", 45), ("Sumber", 12),
    ("Status", 16), ("Catatan", 40),
]
KATEGORI = {"Happy", "Positif", "Negatif", "Edge"}
SUMBER = {"PRD", "PRD-ambigu", "Asumsi"}
ID_RE = re.compile(r"^TC-[A-Z0-9]+-[A-Z0-9]+-\d{3}$")

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
KAT_FILL = {
    "Happy": "E2EFDA", "Positif": "DDEBF7", "Negatif": "FCE4D6", "Edge": "FFF2CC",
}


def validate(name, data):
    errs = []
    seen = set()
    for i, c in enumerate(data.get("cases", []), 1):
        cid = c.get("id", "")
        if not ID_RE.match(cid):
            errs.append(f"case #{i}: id '{cid}' tidak sesuai pola TC-<KODE>-<AKSI>-###")
        if cid in seen:
            errs.append(f"id duplikat: {cid}")
        seen.add(cid)
        if c.get("kategori") not in KATEGORI:
            errs.append(f"{cid}: kategori '{c.get('kategori')}' bukan {sorted(KATEGORI)}")
        if c.get("sumber") not in SUMBER:
            errs.append(f"{cid}: sumber '{c.get('sumber')}' bukan {sorted(SUMBER)}")
    if errs:
        raise SystemExit(f"[{name}] JSON tidak valid:\n  - " + "\n  - ".join(errs))


def fmt_steps(v):
    if isinstance(v, list):
        return "\n".join(f"{i}. {s}" for i, s in enumerate(v, 1))
    return v or ""


def build(path):
    data = json.loads(path.read_text(encoding="utf-8"))
    validate(path.stem, data)
    wb = Workbook()

    ws = wb.active
    ws.title = "Test Cases"
    for col, (title, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "B2"

    for r, c in enumerate(data["cases"], 2):
        values = [
            c["id"], c["kategori"], c.get("skenario", ""), c.get("precondition", ""),
            fmt_steps(c.get("langkah")), c.get("test_data", ""), c.get("expected", ""),
            c["sumber"], c.get("status", "Belum dijalankan"), c.get("catatan", ""),
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.alignment, cell.border = WRAP, BORDER
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=KAT_FILL[c["kategori"]])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(data['cases']) + 1}"

    info = wb.create_sheet("Info")
    rows = [
        ("Modul", data.get("modul", "")),
        ("Kode", f"{data.get('kode', '')}-{data.get('aksi', '')}"),
        ("Halaman", data.get("halaman", "")),
        ("Jumlah TC", len(data["cases"])),
        ("Spec", "\n".join(data["spec"]) if isinstance(data.get("spec"), list) else data.get("spec", "")),
        ("Catatan", data.get("catatan", "")),
    ]
    for n in data.get("notion", []):
        rows.append(("PRD Notion", f"{n.get('judul', '')} — {n.get('url', '')}"))
    for r, (k, v) in enumerate(rows, 1):
        info.cell(row=r, column=1, value=k).font = Font(bold=True)
        info.cell(row=r, column=2, value=v).alignment = WRAP
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 110

    leg = wb.create_sheet("Legenda")
    legend = [
        ("Kolom / Nilai", "Arti"),
        ("Kategori: Happy", "Alur utama yang paling sering dipakai pengguna, data valid."),
        ("Kategori: Positif", "Variasi input/alur valid lain yang harus berhasil."),
        ("Kategori: Negatif", "Input tidak valid / aksi ditolak; sistem harus menolak dengan pesan yang benar."),
        ("Kategori: Edge", "Batas nilai, data kosong, kombinasi tak umum."),
        ("Sumber: PRD", "Expected diambil langsung dari PRD Notion."),
        ("Sumber: PRD-ambigu", "PRD ada tapi tidak tegas; expected = tafsiran, perlu konfirmasi Product."),
        ("Sumber: Asumsi", "Tidak ada di PRD; diturunkan dari perilaku source code / konvensi umum."),
        ("Status: Belum dijalankan", "TC sudah ditulis, belum dieksekusi."),
        ("Status: Pass / Fail / Blocked", "Hasil eksekusi terakhir (diisi setelah run)."),
        ("ID", "TC-<KODE>-<AKSI>-### — sama persis dengan judul `it` di spec Cypress."),
        ("Data uji", "Nama data buatan test diawali 'QA<timestamp>' dan email qa+<ts>@example.invalid."),
    ]
    for r, (k, v) in enumerate(legend, 1):
        a = leg.cell(row=r, column=1, value=k)
        b = leg.cell(row=r, column=2, value=v)
        if r == 1:
            for cell in (a, b):
                cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        b.alignment = WRAP
    leg.column_dimensions["A"].width = 30
    leg.column_dimensions["B"].width = 90

    out = OUT / f"{path.stem}.xlsx"
    wb.save(out)
    return out, len(data["cases"])


def main():
    names = sys.argv[1:]
    files = [SRC / f"{n}.json" for n in names] if names else sorted(SRC.glob("TC_*.json"))
    if not files:
        print("Tidak ada file TC di", SRC)
        return
    for f in files:
        out, n = build(f)
        print(f"OK  {out.relative_to(ROOT)}  ({n} TC)")


if __name__ == "__main__":
    main()
